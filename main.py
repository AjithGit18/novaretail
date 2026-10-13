from pptx import Presentation

ppt_path = "template/novaretail_template.pptx"

prs = Presentation(ppt_path)

for slide_no, slide in enumerate(prs.slides, start=1):
    print(f"\n===== Slide {slide_no} =====")

    for shape in slide.shapes:
        print(shape.name)   